"""CSV and XLSX parsing without NetBox side effects."""

import csv
from pathlib import Path
import re

from atlas.inventory.models import InventoryRow


HEADER_ALIASES = {
    "hostname": "hostname",
    "hardware model": "model",
    "model": "model",
    "management ip": "management_ip",
    "tenant": "tenant",
    "site": "site",
    "device role": "device_role",
    "platform": "platform",
    "vrf": "vrf",
}
REQUIRED = {"hostname", "model", "management_ip"}


def _header(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "").strip().casefold())


def _rows_csv(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        yield from csv.reader(handle)


def _rows_xlsx(path: Path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise ValueError("XLSX support requires the pinned openpyxl dependency") from None
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        for row in sheet.iter_rows(values_only=True):
            yield list(row)
    finally:
        workbook.close()


def parse_inventory(path: Path) -> list[InventoryRow]:
    path = path.expanduser().resolve()
    if not path.is_file():
        raise ValueError(f"Inventory file does not exist: {path}")
    if path.suffix.casefold() == ".csv":
        rows = iter(_rows_csv(path))
    elif path.suffix.casefold() == ".xlsx":
        rows = iter(_rows_xlsx(path))
    else:
        raise ValueError("Inventory file must use .csv or .xlsx")
    try:
        raw_headers = next(rows)
    except StopIteration:
        raise ValueError("Inventory file is empty") from None
    columns: dict[int, str] = {}
    seen = set()
    for index, raw in enumerate(raw_headers):
        name = HEADER_ALIASES.get(_header(raw))
        if name is None:
            continue
        if name in seen:
            raise ValueError(f"Duplicate inventory column: {name}")
        seen.add(name)
        columns[index] = name
    missing = sorted(REQUIRED - seen)
    if missing:
        raise ValueError("Missing required inventory columns: " + ", ".join(missing))

    parsed = []
    for number, values in enumerate(rows, start=2):
        data = {
            name: str(values[index]).strip()
            for index, name in columns.items()
            if index < len(values) and values[index] is not None and str(values[index]).strip()
        }
        if not data:
            continue
        parsed.append(InventoryRow(
            number=number,
            hostname=data.get("hostname", ""),
            model=data.get("model", ""),
            management_ip=data.get("management_ip", ""),
            tenant=data.get("tenant"),
            site=data.get("site"),
            device_role=data.get("device_role"),
            platform=data.get("platform"),
            vrf=data.get("vrf"),
        ))
    return parsed
